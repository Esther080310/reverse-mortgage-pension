"""Extract official CL1 life table from the 2025 experience life table docx.

The source document contains four tables:
0. CL1 pension business table
1. CL2 non-pension business table 1
2. CL3 non-pension business table 2
3. CL4 single-life physical examination table

This project uses CL1 for the liability-side reverse mortgage model.
The table unit is 1/1000, so qx = cell_value / 1000.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document


OUTPUT_CSV = Path("official_cl1_life_table_2025.csv")
OUTPUT_XLSX = Path("official_cl1_life_table_2025.xlsx")
RADIX = 100_000.0
IMPROVEMENT_SCENARIOS = [0.0, 0.15]


def find_source_docx() -> Path:
    candidates = [
        p
        for p in Path(".").glob("*.docx")
        if "生命表" in p.name and "2025" in p.name
    ]
    if not candidates:
        raise FileNotFoundError("Could not find the 2025 life table docx in the workspace.")
    return candidates[0]


def parse_number(text: str) -> float | None:
    text = text.strip().replace(",", "")
    if not text:
        return None
    if not re.fullmatch(r"\d+(\.\d+)?", text):
        return None
    return float(text)


def extract_cl1_base_qx(docx_path: Path) -> pd.DataFrame:
    doc = Document(str(docx_path))
    if len(doc.tables) < 1:
        raise ValueError("The source docx contains no tables.")

    cl1 = doc.tables[0]
    records: list[dict] = []

    for row in cl1.rows[2:]:
        cells = [cell.text.strip() for cell in row.cells]
        for offset in (0, 3, 6):
            if offset + 2 >= len(cells):
                continue
            age_value = parse_number(cells[offset])
            male_value = parse_number(cells[offset + 1])
            female_value = parse_number(cells[offset + 2])
            if age_value is None or male_value is None or female_value is None:
                continue
            age = int(age_value)
            records.append({"sex": "Male", "age": age, "base_qx": male_value / 1000.0})
            records.append({"sex": "Female", "age": age, "base_qx": female_value / 1000.0})

    base = pd.DataFrame(records).drop_duplicates(["sex", "age"]).sort_values(["sex", "age"])
    if base.empty:
        raise ValueError("No CL1 mortality rows were extracted.")
    expected_ages = set(range(int(base["age"].min()), int(base["age"].max()) + 1))
    for sex, sub in base.groupby("sex"):
        missing = expected_ages.difference(set(sub["age"]))
        if missing:
            raise ValueError(f"Missing ages for {sex}: {sorted(missing)[:10]}")
    return base


def build_life_table(base: pd.DataFrame) -> pd.DataFrame:
    all_rows: list[pd.DataFrame] = []
    for improvement in IMPROVEMENT_SCENARIOS:
        for sex, sub in base.groupby("sex"):
            sub = sub.sort_values("age").reset_index(drop=True).copy()
            sub["qx"] = (sub["base_qx"] * (1.0 - improvement)).clip(0.0, 1.0)
            max_age_idx = sub.index[-1]
            sub.loc[max_age_idx, "qx"] = 1.0
            sub["px"] = 1.0 - sub["qx"]

            lx_values = []
            lx = RADIX
            for qx in sub["qx"]:
                lx_values.append(lx)
                lx *= 1.0 - qx
            sub["lx"] = lx_values
            sub["dx"] = sub["lx"] * sub["qx"]
            sub["lx_next"] = sub["lx"] - sub["dx"]
            sub["Lx"] = (sub["lx"] + sub["lx_next"]) / 2.0

            ex_values = []
            future_person_years = sub["Lx"].iloc[::-1].cumsum().iloc[::-1]
            for lx, years in zip(sub["lx"], future_person_years):
                ex_values.append(years / lx if lx > 0 else 0.0)
            sub["ex"] = ex_values
            sub["mortality_improvement"] = improvement

            all_rows.append(
                sub[
                    [
                        "sex",
                        "age",
                        "qx",
                        "px",
                        "lx",
                        "dx",
                        "ex",
                        "mortality_improvement",
                    ]
                ]
            )

    life = pd.concat(all_rows, ignore_index=True)
    return life.sort_values(["mortality_improvement", "sex", "age"]).reset_index(drop=True)


def main() -> None:
    source = find_source_docx()
    base = extract_cl1_base_qx(source)
    life = build_life_table(base)
    life.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    life.to_excel(OUTPUT_XLSX, index=False)

    print(f"Source: {source.resolve()}")
    print(f"Rows: {len(life)}")
    print(f"Ages: {life['age'].min()}-{life['age'].max()}")
    print(f"Scenarios: {sorted(life['mortality_improvement'].unique())}")
    print(f"CSV: {OUTPUT_CSV.resolve()}")
    print(f"XLSX: {OUTPUT_XLSX.resolve()}")


if __name__ == "__main__":
    main()
